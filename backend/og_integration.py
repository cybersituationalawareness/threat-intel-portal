import os
import json
import requests
import re
import uuid
import shutil
import logging

def extract_tlp_color(email_details):
    if not email_details:
        return "Green"

    tlp_pattern = r"\bTLP\s*:\s*(RED|AMBER|GREEN|WHITE|CLEAR)\b"
    match = re.search(tlp_pattern, email_details, re.IGNORECASE)

    if match:
        color = match.group(1).upper()
        if color == 'RED': return "Red"
        if color == 'AMBER': return "Amber"
        if color == 'GREEN': return "Green"
        if color in ['WHITE', 'CLEAR']: return "Clear"

    return "Green" # Default

def extract_text_from_attachment(file_path):
    if not os.path.exists(file_path):
        return None

    ext = os.path.splitext(file_path)[1].lower()
    try:
        if ext in ['.txt', '.log', '.json', '.csv']:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                return f.read().strip()
        elif ext == '.pdf':
            from pypdf import PdfReader
            reader = PdfReader(file_path)
            extracted_text = []
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    extracted_text.append(text)
            return "\n".join(extracted_text).strip()
        elif ext == '.docx':
            import docx
            doc = docx.Document(file_path)
            return "\n".join([p.text for p in doc.paragraphs]).strip()
        elif ext in ['.xlsx', '.xlsm']:
            import openpyxl
            wb = openpyxl.load_workbook(file_path, data_only=True)
            sheet_text = []
            for sheet in wb.worksheets:
                sheet_text.append(f"--- Sheet: {sheet.title} ---")
                for row in sheet.iter_rows(values_only=True):
                    row_vals = [str(cell) for cell in row if cell is not None]
                    if row_vals:
                        sheet_text.append(" | ".join(row_vals))
            return "\n".join(sheet_text).strip()
        else:
            return None
    except Exception as e:
        print(f"[Parser Error]: {e}")
        return None

def get_api_credential():
    # Read from environment variable
    return os.environ.get("GATHER_API_KEY", "")

def get_headers(api_key):
    return {
        "accept": "application/json",
        "Content-Type": "application/json",
        "x-api-key": api_key
    }

def gather_search_cases():
    api_key = get_api_credential()
    if not api_key or api_key == "your_api_key_here":
        print("ERROR: Invalid or missing GATHER_API_KEY")
        return []

    headers = get_headers(api_key)
    url = "https://gather.gov.sg/cms/api/cases/search"
    all_cases = []
    page = 1
    size = 10
    
    while True:
        payload = {
            "statuses" : ["Open"],
            "types" : ["test_ncsc_alert"],
            "order": "desc",
            "sort": "caseRef",
            "page": page,
            "size": size
        }
        
        try:
            response = requests.post(url, headers=headers, json=payload, timeout=10)
            if response.status_code == 200:
                data = response.json().get("data", [])
                if not data:
                    break
                
                all_cases.extend(data)
                
                if len(data) < size:
                    break
                    
                page += 1
            else:
                print(f"Error fetching page {page}: {response.status_code}")
                break
        except Exception as e:
            print(f"Search Network error: {e}")
            break
            
    return all_cases

def gather_get_case_details(case_uuid):
    api_key = get_api_credential()
    headers = get_headers(api_key)
    url = f"https://gather.gov.sg/cms/api/cases/{case_uuid}"
    try:
        response = requests.get(url, headers=headers, timeout=10)
        if response.status_code == 200:
            return response.json()
        return None
    except Exception as e:
        return None

def gather_download_case_attachment(case_uuid, attachment_uuid, save_filename):
    api_key = get_api_credential()
    headers = get_headers(api_key)
    url = f"https://gather.gov.sg/cms/api/cases/{case_uuid}/attachments/{attachment_uuid}"
    try:
        response = requests.get(url, headers=headers, stream=True, timeout=10)
        if response.status_code == 200:
            with open(save_filename, 'wb') as f:
                for chunk in response.iter_content(chunk_size=8192):
                    if chunk:
                        f.write(chunk)
            return True
        return False
    except Exception as e:
        return False

def call_gemini(prompt):
    api_key = os.environ.get("GEMINI_API_KEY")
    if not api_key:
        raise ValueError("GEMINI_API_KEY environment variable is not set.")
    
    url = f"https://generativelanguage.googleapis.com/v1beta/models/gemini-pro:generateContent?key={api_key}"
    headers = {"Content-Type": "application/json"}
    payload = {
        "contents": [{"parts": [{"text": prompt}]}]
    }
    
    response = requests.post(url, headers=headers, json=payload, timeout=30)
    if response.status_code == 200:
        data = response.json()
        return data["candidates"][0]["content"]["parts"][0]["text"].strip()
    else:
        raise Exception(f"Gemini API Error: {response.status_code} - {response.text}")

def interpret_email_with_gemini(email_details):
    if not email_details:
        return "No details provided."

    prompt = f"""You are a strict Cyber Threat Intelligence parser.
       Analyze the following alert details and extract the key points.
       Your response must be valid markdown. Use newline characters to separate each point.
       Do NOT include any conversational filler, introductions, or closing remarks.

       Email Details:
       ---
       {email_details}
       ---

       Respond using the following format, including the markdown for bolding and bullet points:
       **What happened:**
       * [Point 1]

       **Affected systems/users:**
       * [Point 1]

       **Required actions:**
       * [Point 1]"""

    try:
        return call_gemini(prompt)
    except Exception as e:
        print(f"[Gemini Error]: {e}")
        return f"Failed to connect to Gemini API. Raw details: {email_details[:200]}..."

def extract_sla_with_gemini(email_details):
    if not email_details:
        return {"has_sla": False, "sla_value": 0, "sla_unit": "day"}

    prompt = f"""You are an assistant analyzing cyber threat intelligence emails.
Analyze the following email details and determine if there is a Service Level Agreement (SLA) or a required timeframe for action mentioned (e.g. "...within one (1) month...", "within 2 days").
If an SLA is mentioned, extract the numeric value and the unit (must be one of: 'day', 'week', 'month').
Output your result EXACTLY as a valid JSON object with no additional text or markdown formatting. 
Format: {{"has_sla": true or false, "sla_value": integer, "sla_unit": "day" or "week" or "month"}}

Email Details:
---
{email_details}
---
"""
    try:
        text = call_gemini(prompt)
        match = re.search(r'\{.*\}', text, re.DOTALL)
        if match:
            try:
                data = json.loads(match.group(0))
                has_sla = bool(data.get("has_sla", False))
                sla_value = int(data.get("sla_value", 0))
                sla_unit = str(data.get("sla_unit", "day")).lower()
                if sla_unit.endswith('s'):
                    sla_unit = sla_unit[:-1]
                if sla_unit not in ["day", "week", "month"]:
                    sla_unit = "day"
                return {"has_sla": has_sla, "sla_value": sla_value, "sla_unit": sla_unit}
            except (json.JSONDecodeError, ValueError):
                pass
    except Exception as e:
        print(f"[Gemini SLA Error]: {e}")

    return {"has_sla": False, "sla_value": 0, "sla_unit": "day"}

def categorize_email_with_gemini(email_details):
    if not email_details:
        return "Other"

    prompt = f"""You are an expert Cyber Threat Intelligence analyst.
Analyze the following email details and categorize the alert or advisory into ONE most appropriate category from the following examples (or a similar concise category):
- Exploited Vulnerabilities
- Indicators of Compromise
- Campaign
- Threat Hunt Package
- RFI
- Other

Respond ONLY with the category name (e.g. Exploited Vulnerabilities) without any punctuation, markdown, prefix, explanation, or quotes.

Email Details:
---
{email_details}
---
"""
    try:
        category = call_gemini(prompt)
        category = category.replace('"', '').replace("'", "").replace('**', '').strip()
        category = category.split('\n')[0].strip()
        if len(category) > 50 or not category:
            return "General Threat Intelligence"
        return category
    except Exception as e:
        print(f"[Gemini Categorize Error]: {e}")

    return "General Threat Intelligence"

def sync_latest_intel_from_og(existing_case_ids=None):
    """
    Main function called by FastAPI. Returns a list of dicts that can be passed 
    to ThreatIntelCreate schema.
    """
    search_results = gather_search_cases()
    if not search_results:
        return []

    parsed_intels = []
    if existing_case_ids is None:
        existing_case_ids = set()
    
    for case in search_results:
        case_uuid = case.get('uuid')
        case_ref = case.get('caseRef', f'OG-{uuid.uuid4().hex[:6]}')
        
        if case_ref in existing_case_ids:
            continue
            
        case_details = gather_get_case_details(case_uuid)
        if not case_details:
            continue

        case_data = case_details.get("data", {})
        fields = case_data.get('fields', {})
        email_from = str(fields.get('From', ''))
        title = str(fields.get('Subject', 'Untitled Alert'))
        email_details = str(fields.get('Details', ''))

        # Classify Type
        intel_type = 'Alert'
        if 'advisory' in title.lower():
            intel_type = 'Advisory'

        # Classify Scope/Tags
        source = 'National'
        confidence = 'Medium'
        target_member = 'All Member'
        tags = []
        if 'ncsc' in email_from.lower() or 'ncsc' in title.lower():
            source = 'National'
            target_member = 'All Member'
            tags.append('NCSC')
            confidence = 'High'

        elif 'caas.gov.sg' in email_from.lower() or 'caas' in title.lower():
            source = 'Threat Intel'
            target_member = 'All Member'
            tags.append('Threat Intel')
            confidence = 'Medium'




        tlp_color = extract_tlp_color(email_details)
        sla_data = extract_sla_with_gemini(email_details)
        category = categorize_email_with_gemini(email_details)
        
        # Parse description with Gemini and fix clustered formatting
        description = interpret_email_with_gemini(email_details)
        # Convert bullet points from '* ' to '- '
        description = re.sub(r'(?<!\*)\* ', '\n- ', description)
        description = re.sub(r'(\*\*[^*]+\*\*)', r'\n\n\1\n', description)
        description = re.sub(r'[ \t]+\n', '\n', description)
        description = re.sub(r'\n[ \t]+', '\n', description)
        # Remove extra blank lines between bullet points
        description = re.sub(r'\n{2,}- ', '\n- ', description)
        description = re.sub(r'\n{3,}', '\n\n', description)
        description = description.strip()

        # Handle Attachments for Threat Data
        threat_data = "No attachment data available."
        attachments_data = case_data.get("attachments")
        downloaded_files = []

        if isinstance(attachments_data, list):
            for att in attachments_data:
                real_uuid = att.get("uuid")
                fname = att.get("name", f"download_{uuid.uuid4().hex[:6]}.dat")
                if gather_download_case_attachment(case_uuid, real_uuid, fname):
                    downloaded_files.append(fname)
        elif isinstance(attachments_data, dict):
            for key, att in attachments_data.items():
                if isinstance(att, dict):
                    real_uuid = att.get("uuid", key)
                    fname = att.get("name", f"download_{uuid.uuid4().hex[:6]}.dat")
                    if gather_download_case_attachment(case_uuid, real_uuid, fname):
                        downloaded_files.append(fname)

        if downloaded_files:
            # For threat_data extraction, just use the first file for now
            attachment_text = extract_text_from_attachment(downloaded_files[0])
            if attachment_text:
                threat_data = attachment_text[:2000] # Limit size for DB
                # Remove defang characters
                threat_data = threat_data.replace("[", "").replace("]", "")
            
        parsed_intels.append({
            "type": intel_type,
            "title": title,
            "case_id": case_ref,
            "description": description,
            "threat_data": threat_data,
            "tlp": tlp_color,
            "confidence": confidence,
            "tags": tags,
            "classification": source,
            "category": category,
            "has_sla": sla_data["has_sla"],
            "sla_value": sla_data["sla_value"],
            "sla_unit": sla_data["sla_unit"],
            "temp_attachments": downloaded_files
        })

    return parsed_intels
